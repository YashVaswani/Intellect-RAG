# ingestion/pptx_loader.py
"""
PowerPoint (.pptx) text extraction.
Extracts text from slides, notes, and shapes.
"""
import logging

logger = logging.getLogger(__name__)


def load(file_path: str) -> list[dict]:
    """
    Extract text from a PowerPoint presentation.
    Each slide becomes a segment.
    
    Returns:
        List of {"text": str, "page_number": int, "metadata": dict}
    """
    try:
        from pptx import Presentation

        prs = Presentation(file_path)
        segments = []

        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_texts = []

            # Extract text from all shapes
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_texts.append(text)

                # Extract from tables in slides
                if shape.has_table:
                    for row in shape.table.rows:
                        row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_data:
                            slide_texts.append(" | ".join(row_data))

            # Extract speaker notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_texts.append(f"Speaker Notes: {notes}")

            if slide_texts:
                segments.append({
                    "text": "\n".join(slide_texts),
                    "page_number": slide_num,
                    "metadata": {"extraction_method": "pptx", "slide_number": slide_num},
                })

        logger.info(f"PPTX loaded: {file_path} → {len(segments)} slides")
        return segments

    except ImportError:
        logger.error("python-pptx not installed. Run: pip install python-pptx")
        return []
    except Exception as e:
        logger.error(f"Failed to load PPTX {file_path}: {e}")
        return []
